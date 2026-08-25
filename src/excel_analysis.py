from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

_SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm")

_LABEL_KEYWORDS = {
    "objectives": ("objective", "goal"),
    "scope": ("scope",),
    "kpis": ("kpi", "metric"),
    "filters": ("filter",),
    "calculations": ("calculation", "formula"),
    "visual_requirements": ("visual", "chart", "graph", "dashboard"),
    "reporting_expectations": ("reporting expectation", "expectation", "frequency"),
}

_HEADER_KPI_KEYWORDS = (
    "revenue", "sales", "profit", "total", "count", "rate", "percent", "%",
    "kpi", "metric", "amount", "cost", "margin",
)

_HEADER_FILTER_KEYWORDS = (
    "region", "date", "status", "category", "filter", "segment",
    "department", "quarter", "year", "month",
)


class UnsupportedExcelFormatError(Exception):
    """Raised when the file extension is not a supported Excel format."""


class CorruptExcelFileError(Exception):
    """Raised when the file cannot be opened or parsed as a workbook."""


def analyze_report(path: str) -> dict:
    """Parse an Excel report and extract objectives, scope, KPIs, filters,
    calculations, visual requirements, and reporting expectations.

    Examines worksheet names (scope), header rows of tabular sheets (KPIs,
    filters), formula cells (calculations), and free-text "Label: value"
    cells anywhere in the workbook (all categories).
    """
    if not path.lower().endswith(_SUPPORTED_EXTENSIONS):
        raise UnsupportedExcelFormatError(f"Unsupported file format: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (InvalidFileException, BadZipFile, OSError) as exc:
        raise CorruptExcelFileError(f"Could not open Excel file: {path}") from exc

    result: dict = {category: [] for category in _LABEL_KEYWORDS}

    for sheet in workbook.worksheets:
        _add_unique(result["scope"], sheet.title)
        _scan_headers(sheet, result)
        _scan_cells(sheet, result)

    workbook.close()
    return result


def _add_unique(items: list, value: str) -> None:
    value = value.strip()
    if value and value not in items:
        items.append(value)


def _scan_headers(sheet, result: dict) -> None:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
    if len(header_row) < 2:
        return

    for cell in header_row:
        text = str(cell.value).strip() if cell.value is not None else ""
        if not text:
            continue
        lowered = text.lower()
        if any(keyword in lowered for keyword in _HEADER_KPI_KEYWORDS):
            _add_unique(result["kpis"], text)
        if any(keyword in lowered for keyword in _HEADER_FILTER_KEYWORDS):
            _add_unique(result["filters"], text)


def _scan_cells(sheet, result: dict) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if not text:
                continue
            if text.startswith("="):
                _add_unique(result["calculations"], text.lstrip("="))
            else:
                _match_label(text, result)


def _match_label(text: str, result: dict) -> None:
    lowered = text.lower()
    label, _, value = text.partition(":")
    for category, keywords in _LABEL_KEYWORDS.items():
        if not any(keyword in lowered for keyword in keywords):
            continue
        if value and any(keyword in label.lower() for keyword in keywords):
            _add_unique(result[category], value.strip())
            return
        if not value and category in ("visual_requirements", "reporting_expectations"):
            _add_unique(result[category], text)
            return
